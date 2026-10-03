import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="none", right="none"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="{"single" if top != "none" else "none"}" w:sz="4" w:space="0" w:color="{top if top != "none" else "auto"}"/>
            <w:left w:val="{"single" if left != "none" else "none"}" w:sz="4" w:space="0" w:color="{left if left != "none" else "auto"}"/>
            <w:bottom w:val="{"single" if bottom != "none" else "none"}" w:sz="4" w:space="0" w:color="{bottom if bottom != "none" else "auto"}"/>
            <w:right w:val="{"single" if right != "none" else "none"}" w:sz="4" w:space="0" w:color="{right if right != "none" else "auto"}"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def add_placeholder_box(doc, fig_number, caption_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.2)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=360, bottom=360, left=240, right=240)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="94A3B8"/>
            <w:left w:val="single" w:sz="8" w:space="0" w:color="94A3B8"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="94A3B8"/>
            <w:right w:val="single" w:sz="8" w:space="0" w:color="94A3B8"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("[INSERT SCREENSHOT HERE]")
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(71, 85, 105) # Slate 600
    
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(14)
    run_cap = cap.add_run(f"Figure {fig_number}: {caption_text}")
    run_cap.font.name = "Calibri"
    run_cap.font.size = Pt(10)
    run_cap.font.bold = True
    run_cap.font.italic = True
    run_cap.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

def create_document():
    doc = Document()
    
    # Page setup - A4
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("AI-Based Sentiment Analysis System | Academic Project Submission")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun_left = fp.add_run("Candidate: Harshitha Siddagattu")
        frun_left.font.name = "Calibri"
        frun_left.font.size = Pt(8.5)
        frun_left.font.color.rgb = RGBColor(100, 116, 139)
        
    # Styles
    navy = RGBColor(30, 58, 138) # #1E3A8A
    dark = RGBColor(15, 23, 42)   # #0F172A
    muted = RGBColor(71, 85, 105) # #475569
    
    def add_section_heading(number, title):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(f"{number}. {title.upper()}")
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = navy
        
        # Bottom rule under heading
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="1E3A8A"/></w:pBdr>')
        h._p.get_or_add_pPr().append(pBdr)

    def add_para(text, bold_prefix="", space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = dark
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = dark
        return p

    # ==================== 1. PROJECT TITLE ====================
    add_section_heading(1, "Project Title")
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("AI-Based Sentiment Analysis System")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = dark

    # ==================== 2. CANDIDATE NAME(S) ====================
    add_section_heading(2, "Candidate Name(s)")
    add_para("Harshitha Siddagattu", bold_prefix="Candidate Name(s): ")

    # ==================== 3. DOMAIN ====================
    add_section_heading(3, "Domain")
    add_para("Artificial Intelligence / Machine Learning / Natural Language Processing")

    # ==================== 4. DATASET ====================
    add_section_heading(4, "Dataset")
    add_para("A comprehensive description of the datasets used for development, validation, and empirical benchmarking of the sentiment intelligence platform is presented below:")

    # Table for dataset specifications
    ds_table = doc.add_table(rows=10, cols=2)
    ds_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    col_widths = [Inches(2.2), Inches(4.0)]
    for row in ds_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    dataset_rows = [
        ("Dataset Name", "SentinelAI Omnichannel Customer Feedback & Evaluation Benchmark"),
        ("Dataset Source", "Verified customer feedback across digital touchpoints (Website Reviews, Google Reviews, Customer Surveys, Email Support Tickets, App Store Reviews)"),
        ("Dataset URL", "Local files: data/sample_feedback.csv & data/evaluation_sample.csv\n(Repository: https://github.com/JwHarshitha/AI_Projects_Repo)"),
        ("Dataset Type", "Tabular Text / Comma-Separated Values (.csv)"),
        ("Number of Records", "160 total verified records (120 records in sample_feedback.csv + 40 records in evaluation_sample.csv)"),
        ("Features / Columns", "• sample_feedback.csv: text, date, source, category\n• evaluation_sample.csv: text, label"),
        ("Target Variable", "Sentiment Polarity (sentiment / label)"),
        ("Sentiment Classes", "Three (3) Classes: positive, neutral, negative"),
        ("Category Breakdown", "Product Quality, Delivery, Customer Support, Pricing, Mobile App, Website"),
        ("Dataset Description", "A curated, balanced dataset specifically designed to assess multi-aspect customer reviews, negation phrasing (e.g., 'not good', 'not bad'), valence shifters, and Plutchik emotional states without artificial repetitiveness.")
    ]

    for idx, (label, val) in enumerate(dataset_rows):
        c0 = ds_table.cell(idx, 0)
        c1 = ds_table.cell(idx, 1)
        
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
        set_cell_margins(c1, top=100, bottom=100, left=140, right=140)
        set_cell_borders(c0, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1")
        set_cell_borders(c1, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = navy
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = dark

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ==================== 5. SOURCE CODE ====================
    add_section_heading(5, "Source Code")
    add_para("The system is implemented as a production-grade, full-stack sentiment intelligence web application consisting of a modular Python/FastAPI backend and a React/TypeScript analytics frontend.")
    
    add_para("React 18, Vite, TypeScript, Tailwind CSS, Recharts, Lucide Icons, and Framer Motion located under frontend/src/.", bold_prefix="• Frontend Source Code: ")
    add_para("Python 3.9+, FastAPI, SQLAlchemy ORM, and Uvicorn ASGI server located under backend/app/.", bold_prefix="• Backend Source Code: ")
    add_para("scikit-learn, NumPy, Pandas, N-gram TF-IDF vectorization, valence shifting rules, clause-level Aspect-Based Sentiment Analysis (ABSA), and Plutchik/Ekman emotion classification located under backend/app/ml/ and backend/app/services/.", bold_prefix="• Machine Learning Source Code: ")
    add_para("https://github.com/JwHarshitha/AI_Projects_Repo", bold_prefix="• Source Code Repository: ")

    add_para("Project Folder Structure:", bold_prefix="Project Structure: ")
    
    # Folder Structure Box
    fs_table = doc.add_table(rows=1, cols=1)
    fs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fs = fs_table.cell(0, 0)
    c_fs.width = Inches(6.2)
    set_cell_background(c_fs, "0F172A") # Near black
    set_cell_margins(c_fs, top=140, bottom=140, left=160, right=160)
    set_cell_borders(c_fs, top="1E293B", bottom="1E293B", left="1E293B", right="1E293B")
    
    tree_text = (
        "AI_Projects_Repo/\n"
        "├── backend/\n"
        "│   ├── app/\n"
        "│   │   ├── api/             # 11 REST API Routers (auth, analyze, datasets, feedback, etc.)\n"
        "│   │   ├── ml/              # Base model interface, TF-IDF classifier, Lexicon rules\n"
        "│   │   ├── models/          # SQLAlchemy Database entities (User, FeedbackRecord, etc.)\n"
        "│   │   ├── schemas/         # Pydantic v2 validation schemas\n"
        "│   │   ├── services/        # ABSA, Emotion, Keyword, Insight, Evaluation services\n"
        "│   │   ├── utils/           # Database auto-seeder & demo loader\n"
        "│   │   ├── config.py        # Centralized settings & configurations\n"
        "│   │   ├── database.py      # Database engine & SessionLocal factory\n"
        "│   │   └── main.py          # FastAPI application & lifecycle mounting\n"
        "│   ├── requirements.txt     # Python backend dependencies\n"
        "│   ├── run.py               # Uvicorn server entrypoint script\n"
        "│   └── sentinel.db          # SQLite local persistent database\n"
        "├── frontend/\n"
        "│   ├── src/\n"
        "│   │   ├── components/      # StatCards, SentimentBadge, ConfidenceMeter, Drawers\n"
        "│   │   ├── context/         # AuthContext & state management\n"
        "│   │   ├── pages/           # 11 major production screens\n"
        "│   │   ├── services/        # Typed Axios/Fetch API client\n"
        "│   │   ├── types/           # TypeScript interfaces\n"
        "│   │   ├── App.tsx          # Client router & command palette shell\n"
        "│   │   └── main.tsx         # React root mounting\n"
        "│   ├── package.json         # Node.js dependencies\n"
        "│   └── vite.config.ts       # Vite proxy and build configuration\n"
        "├── data/\n"
        "│   ├── sample_feedback.csv  # 120 verified customer feedback records\n"
        "│   └── evaluation_sample.csv # 40 labeled benchmark records\n"
        "├── docker/                  # Dockerfile.backend & Dockerfile.frontend\n"
        "├── start.sh                 # Single-command fullstack startup script\n"
        "└── README.md                # Comprehensive system documentation"
    )
    p_tree = c_fs.paragraphs[0]
    p_tree.paragraph_format.space_before = Pt(0)
    p_tree.paragraph_format.space_after = Pt(0)
    r_tree = p_tree.add_run(tree_text)
    r_tree.font.name = "Courier New"
    r_tree.font.size = Pt(8.0)
    r_tree.font.color.rgb = RGBColor(226, 232, 240) # Slate 200

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ==================== 6. JUPYTER NOTEBOOK ====================
    add_section_heading(6, "Jupyter Notebook")
    add_para("Jupyter Notebook: To be attached")

    # ==================== 7. MODEL FILES ====================
    add_section_heading(7, "Model Files")
    add_para("Information regarding the trained and serialized Machine Learning model pipeline is provided below:")

    model_table = doc.add_table(rows=6, cols=2)
    model_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in model_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    model_rows = [
        ("Model Name", "Sentinel Calibrated Sentiment Classification Pipeline"),
        ("Model Type", "scikit-learn Pipeline (TfidfVectorizer + L2-Regularized Logistic Regression)"),
        ("Purpose", "Inference of 3-class sentiment polarity (positive, neutral, negative) with calibrated confidence scoring and clause-level aspect extraction"),
        ("File Name", "backend/app/ml/sentiment_pipeline.joblib"),
        ("Preprocessing Files", "In-pipeline tokenization, lowercase normalization, and sublinear term frequency weighting embedded in pipeline"),
        ("Label Encoder / Vectorizer Files", "TfidfVectorizer(ngram_range=(1, 2), sublinear_tf=True) serialized directly inside sentiment_pipeline.joblib")
    ]

    for idx, (label, val) in enumerate(model_rows):
        c0 = model_table.cell(idx, 0)
        c1 = model_table.cell(idx, 1)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
        set_cell_margins(c1, top=100, bottom=100, left=140, right=140)
        set_cell_borders(c0, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1")
        set_cell_borders(c1, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = navy
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = dark

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ==================== 8. APPLICATION / DASHBOARD ====================
    add_section_heading(8, "Application / Dashboard")
    add_para("SentinelAI – AI Sentiment Intelligence Platform", bold_prefix="• Application Name: ")
    add_para("A commercial-grade AI intelligence platform designed to ingest, process, and analyze multi-channel customer reviews, feedback, and support interactions to derive actionable operational insights, emotional dynamics, and aspect-level drivers.", bold_prefix="• Purpose: ")
    
    add_para("Key Architectural Features:", bold_prefix="• Main Features: ")
    features = [
        ("Sentiment Classification: ", "High-precision 3-class polarity estimation with calibrated probability distributions."),
        ("Emotion Detection: ", "Quantification of Plutchik and Ekman emotions (Joy, Trust, Anger, Sadness, Fear, Surprise, Disgust)."),
        ("Aspect-Based Sentiment Analysis (ABSA): ", "Multi-clause syntax parsing to isolate sentiment across Product Quality, Shipping, Support, Pricing, Mobile App, and Web."),
        ("Explainable AI ('Analysis Evidence'): ", "Exact linguistic attribution highlighting positive signals, negative friction cues, and intensifiers."),
        ("Batch CSV Ingestion: ", "Asynchronous background processing with live progress reporting and automatic column detection."),
        ("Interactive Feedback Explorer: ", "Searchable tabular interface with multi-dimensional filtering, sorting, pagination, and slide-over drawer."),
        ("Automated AI Insights & Executive Summary: ", "Statistically grounded business insights (Positive Momentum, Negative Signal, Emerging Topics) and executive synthesis."),
        ("Model Benchmarking Hub: ", "Live evaluation dataset upload computing empirical Accuracy, Precision, Recall, F1 Score, and Confusion Matrix."),
        ("Report Generation & Export: ", "Automated generation of executive reports with print-ready PDF styling and CSV dataset exports.")
    ]
    for title, descr in features:
        add_para(descr, bold_prefix=f"  - {title}", space_after=3)

    add_para("Frontend (React 18, Vite, TypeScript, Tailwind CSS, Recharts) | Backend (Python, FastAPI, Uvicorn) | Database (SQLite via SQLAlchemy ORM) | Machine Learning (Calibrated TF-IDF + Logistic Regression, NLTK/Lexicon Valence Rules).", bold_prefix="• Technology Stack: ")
    
    add_para("Major Application Modules:", bold_prefix="• Application Modules: ")
    modules = [
        ("Dashboard: ", "High-level KPI overview, sentiment donut chart, 14-day trend line, emotion horizontal bar, and operational topic cards."),
        ("Single Analysis: ", "Real-time deep NLP test workbench with interactive clause breakdown, confidence meters, and linguistic evidence."),
        ("Batch Analysis: ", "Drag-and-drop CSV processor supporting asynchronous background execution and category aggregations."),
        ("Feedback Explorer: ", "Full data table with keyword search, multi-variable filters, and diagnostic detail drawer."),
        ("Sentiment Trends: ", "Longitudinal stacked area graphs, emotion trajectory tracking, and statistical anomaly detection alerts."),
        ("Model Performance: ", "Architecture comparison cards, benchmark visualizer, and empirical dataset evaluation utility."),
        ("Reports: ", "Executive dossier generator with dataset scoping, thematic synthesis, print formatting, and CSV data export.")
    ]
    for m_title, m_desc in modules:
        add_para(m_desc, bold_prefix=f"  - {m_title}", space_after=3)

    add_para("• Local / Deployed URL:")
    add_para("http://localhost:5173", bold_prefix="  - Frontend Application URL: ")
    add_para("http://127.0.0.1:8000", bold_prefix="  - Backend REST API URL: ")
    add_para("http://127.0.0.1:8000/docs", bold_prefix="  - Interactive Swagger API Documentation: ")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ==================== 9. SCREENSHOTS ====================
    add_section_heading(9, "Screenshots")
    add_para("The following figures represent the primary user interface screens of the SentinelAI sentiment intelligence platform:")

    add_placeholder_box(doc, 1, "Application Dashboard")
    add_placeholder_box(doc, 2, "Sentiment Analysis Page")
    add_placeholder_box(doc, 3, "Sentiment Analysis Result")
    add_placeholder_box(doc, 4, "Batch Analysis")
    add_placeholder_box(doc, 5, "Feedback Explorer")
    add_placeholder_box(doc, 6, "Model Performance")
    add_placeholder_box(doc, 7, "Reports / Analytics")

    # ==================== 10. DOCUMENTATION ====================
    add_section_heading(10, "Documentation")
    add_para("Comprehensive project documentation detailing system architecture, mathematical methodologies, API specifications, and operational guides is maintained in the repository:")

    doc_topics = [
        ("Project Overview: ", "Comprehensive operational vision of SentinelAI as an enterprise sentiment intelligence platform."),
        ("System Architecture: ", "Decoupled client-server design leveraging FastAPI REST endpoints, asynchronous task processing, and React component modularity."),
        ("Dataset Documentation: ", "Complete data dictionaries, sample distributions, and validation sets under data/sample_feedback.csv and data/evaluation_sample.csv."),
        ("Data Preprocessing: ", "Text normalization, stop-word filtering, n-gram TF-IDF matrix generation, and sublinear scaling."),
        ("Machine Learning Methodology: ", "Hybrid NLP combining regularized linear classification with valence shifting windows (negation handling and intensifier scaling)."),
        ("Model Documentation: ", "Theoretical formulations, calibration parameters, hyperparameter optimization, and empirical benchmark validation results."),
        ("Application Documentation: ", "Module-by-module feature walkthroughs, UI design tokens, and state management lifecycle."),
        ("API Documentation: ", "OpenAPI / Swagger documentation available interactively at /docs across 11 functional REST routers."),
        ("Installation Instructions: ", "Step-by-step setup guide for Python virtual environments, Node.js packages, and Docker Compose deployment."),
        ("User Guide: ", "End-user workflows for ad-hoc single text analysis, bulk CSV ingestion, report generation, and CSV exporting."),
        ("Testing Documentation: ", "Unit testing, empirical classification metric evaluation (Accuracy, Precision, Recall, F1), and API endpoint verification."),
        ("Future Scope: ", "Integration with cloud GPU transformer microservices (DistilRoBERTa), multilingual cross-lingual models, and automated customer service webhook triggers.")
    ]
    for d_title, d_desc in doc_topics:
        add_para(d_desc, bold_prefix=f"• {d_title}", space_after=3)

    add_para("README.md (Located in root repository directory)", bold_prefix="• Primary Documentation File: ")

    # ==================== 11. GITHUB REPOSITORY ====================
    add_section_heading(11, "GitHub Repository")
    add_para("The complete project codebase, machine learning pipeline, dataset files, frontend components, and launch scripts are hosted at the following GitHub repository:")
    add_para("https://github.com/JwHarshitha/AI_Projects_Repo", bold_prefix="GitHub Repository: ")

    # Save document
    output_filename = "AI_Based_Sentiment_Analysis_System_Project_Document.docx"
    doc.save(output_filename)
    print(f"Document successfully created: {output_filename}")

if __name__ == "__main__":
    create_document()
